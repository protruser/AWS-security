import hmac
import io
import json
import os
import threading
from datetime import datetime, timedelta, timezone
from functools import wraps
from xml.sax.saxutils import escape as _xml_escape

import boto3
from botocore.config import Config
from dotenv import load_dotenv
from flask import Flask, jsonify, request, send_file, session
from flask_cors import CORS
from openai import OpenAI
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from db import get_connection
from services.event_analysis_service import AnalysisBusy, EventNotFound, analyze_event
from services.ai_diagnosis_service import (
    CATEGORY_ORDER,
    DiagnosisError,
    diagnose_aws_state,
    load_rule_meta,
)
from services.aws_collector import collect_aws_state

load_dotenv()

# 배포용 컨테이너에서는 프론트엔드 빌드 결과(frontend/dist)를 이 Flask가 함께 서빙한다.
# 로컬 개발(vite dev)에서는 이 폴더가 없어도 API 서버로는 정상 동작한다.
app = Flask(__name__, static_folder="static", static_url_path="")
app.secret_key = os.getenv("FLASK_SECRET_KEY", "dev-only-change-this-secret")
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "false").lower() == "true",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
)

origins = [
    origin.strip()
    for origin in os.getenv(
        "CORS_ORIGINS", "http://localhost:8443,http://localhost:5173"
    ).split(",")
    if origin.strip()
]
CORS(
    app,
    resources={r"/api/*": {"origins": origins}},
    supports_credentials=True,
)

VALID_SEVERITIES = {"Critical", "High", "Medium", "Low", "Info"}
REMEDIATION_ACTIONS = {
    "sqli": "block_ip", "dir": "block_ip", "brute": "block_ip",
    "xss": "block_ip", "cred": "disable_access_key", "port": "block_ip",
}
REMEDIATION_CLOSED_STATUSES = {"조치 완료", "자동 완료", "완료", "예외 처리"}
LOG_RANGE_DELTAS = {
    "15m": timedelta(minutes=15),
    "1h": timedelta(hours=1),
    "6h": timedelta(hours=6),
    "24h": timedelta(hours=24),
    "7d": timedelta(days=7),
}
LOG_SOURCES = {"waf", "guardduty", "inspector"}
# 운영 지표는 infra의 Lambda C 가 5분마다 채우는 service_metrics 표에서 읽는다
# (Terraform: modules/lambda_c, modules/lambda_common/common/db.py).
OVERVIEW_SERIES_LIMIT = 20
MONITORING_METRIC_ALIASES = {
    "all": "all",
    "cpu": "cpu",
    "memory": "memory",
    "latency": "latency",
    "rps": "rps",
    "errorrate": "errorRate",
    "error-rate": "errorRate",
    "error_rate": "errorRate",
    "health": "health",
}


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("authenticated"):
            return jsonify({"error": "UNAUTHORIZED", "message": "로그인이 필요합니다."}), 401
        return view(*args, **kwargs)

    return wrapped


def role_required(*allowed_roles):
    """login_required에 역할 제한을 추가한다. 관리자는 조치 요청을 보낼 수
    있지만 직접 승인할 수 없고, 승인자는 그 반대다 - 역할이 섞이면 요청자와
    승인자가 같은 사람이 되어버려서 승인 절차 자체가 의미 없어진다."""

    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not session.get("authenticated"):
                return jsonify({"error": "UNAUTHORIZED", "message": "로그인이 필요합니다."}), 401
            if session.get("role") not in allowed_roles:
                return jsonify({"error": "FORBIDDEN", "message": "이 작업을 수행할 권한이 없습니다."}), 403
            return view(*args, **kwargs)

        return wrapped

    return decorator


def _current_user():
    return {
        "username": session.get("username", "admin"),
        "role": session.get("role", os.getenv("ADMIN_ROLE", "관리자")),
        "team": session.get("team", os.getenv("ADMIN_TEAM", "보안관제팀")),
    }


def _json_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            return parsed if isinstance(parsed, list) else []
        except json.JSONDecodeError:
            return []
    return []


KST = timezone(timedelta(hours=9))


def _open_group_ids(cursor, scenario_type, attacker_ip, title):
    """대시보드 목록 조회(_read_dashboard_data)가 하나로 합쳐서 보여주는 것과 정확히
    같은 기준(scenario_type + 공격자 IP가 있으면 IP, 없으면 title)으로, 아직 안 닫힌
    이벤트 id를 전부 찾는다. 예외 처리/조치 승인이 "화면에 보이는 대표 1건"만 닫으면
    나머지가 다음 새로고침 때 새 대표로 다시 튀어나오므로, 항상 그룹 전체를 같이
    처리해야 한다."""
    if attacker_ip:
        cursor.execute(
            "SELECT id FROM security_events WHERE status NOT IN "
            "('조치 완료', '자동 완료', '예외 처리', '완료') "
            "AND scenario_type = %s AND attacker_ip = %s",
            (scenario_type, attacker_ip),
        )
    else:
        cursor.execute(
            "SELECT id FROM security_events WHERE status NOT IN "
            "('조치 완료', '자동 완료', '예외 처리', '완료') "
            "AND scenario_type = %s AND title = %s",
            (scenario_type, title),
        )
    return [row["id"] for row in cursor.fetchall()]


def _format_datetime(value, fmt="%Y.%m.%d %H:%M"):
    """DB의 시각은 항상 UTC(UTC_TIMESTAMP())로 저장돼 있다. 그대로 strftime하면
    한국 시간보다 9시간 느리게 표시되므로, 화면에 보여줄 때는 KST로 바꿔서 찍는다."""
    if value is None:
        return "-"
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc).astimezone(KST).strftime(fmt)
    return str(value)


def _elapsed_text(value):
    if not isinstance(value, datetime):
        return "-"

    now = datetime.now(value.tzinfo or timezone.utc)
    if value.tzinfo is None:
        now = datetime.now()

    minutes = max(0, int((now - value).total_seconds() // 60))
    if minutes < 60:
        return f"{minutes}분"

    hours, remain = divmod(minutes, 60)
    if hours < 24:
        return f"{hours}시간 {remain}분"

    days, remain_hours = divmod(hours, 24)
    return f"{days}일 {remain_hours}시간"


def _severity(value):
    text = str(value or "Info").capitalize()
    return text if text in VALID_SEVERITIES else "Info"


def _event_to_action_event(row):
    return {
        "id": str(row["id"]),
        "severity": _severity(row.get("severity")),
        "title": row.get("title") or "보안 이벤트",
        "service": row.get("service") or "Unknown",
        "scenarioType": row.get("scenario_type") or "",
        "asset": row.get("asset") or "-",
        "detectedAt": _format_datetime(row.get("detected_at")),
        "elapsed": _elapsed_text(row.get("detected_at")),
        # 같은 종류(scenario_type)로 같은 공격자 IP(없으면 같은 제목)의 반복 탐지는
        # 목록 조회 시 최신 것 하나로 합쳐서 보여준다. 이 값이 2 이상이면 "N번 반복
        # 감지됨"이라는 뜻 - 실제로 그만큼의 별도 행이 DB에 있다.
        "occurrenceCount": int(row.get("occurrence_count") or 1),
        "status": row.get("status") or "검토 필요",
        "recommendation": row.get("recommendation") or "",
        "autoRemediation": bool(row.get("auto_remediation")),
        "remediationType": "AUTO" if row.get("scenario_type") in REMEDIATION_ACTIONS else "MANUAL",
        "highlightAssets": _json_list(row.get("highlight_assets")),
        "attackPath": _json_list(row.get("attack_path")),
        "details": {
            "attackerIP": row.get("attacker_ip") or "",
            "requestURL": row.get("request_url") or "",
            "rule": row.get("rule_name") or "",
            "blocked": (
                bool(row.get("blocked")) if row.get("blocked") is not None else None
            ),
            "logs": row.get("logs") or "",
        },
    }


def _event_to_detect_history(row):
    if row.get("block_result"):
        blocked_text = row["block_result"]
    elif row.get("blocked") is True or row.get("blocked") == 1:
        blocked_text = "차단"
    else:
        blocked_text = "-"

    detected_at = row.get("detected_at")
    return {
        "id": str(row["id"]),
        "time": _format_datetime(detected_at, "%H:%M"),
        "sev": _severity(row.get("severity")),
        "event": row.get("title") or "보안 이벤트",
        "service": row.get("service") or "Unknown",
        "asset": row.get("asset") or "-",
        "ip": row.get("attacker_ip") or "-",
        "blocked": blocked_text,
        "status": row.get("status") or "-",
    }


def _remediation_to_history(row):
    completed = row.get("completed_at")
    requested = row.get("requested_at")
    return {
        "id": str(row["id"]),
        "time": _format_datetime(requested or completed, "%H:%M"),
        "event": row.get("event_title") or row.get("action_type") or "보안 조치",
        "asset": row.get("asset") or "-",
        "ip": row.get("attacker_ip") or None,
        "method": row.get("method") or "수동",
        "approver": row.get("approver") or "-",
        "result": row.get("result") or row.get("status") or "-",
        "completedAt": _format_datetime(completed, "%H:%M:%S"),
        "occurrenceCount": int(row.get("occurrence_count") or 1),
    }


def _event_to_log(row):
    detected_at = row.get("detected_at")
    return {
        "id": str(row["id"]),
        "service": row.get("service") or "-",
        "scenarioType": row.get("scenario_type") or "-",
        "severity": _severity(row.get("severity")),
        "title": row.get("title") or "보안 이벤트",
        "asset": row.get("asset") or None,
        # DB에 저장된 시각은 항상 UTC라, 타임존 표기 없이 그냥 isoformat()만 하면
        # 브라우저(JS)가 이걸 "이미 로컬시간"으로 착각해서 9시간 밀려 보인다.
        # UTC임을 명시해서 내려주면 프론트에서 알아서 KST로 변환해 표시한다.
        "detectedAt": (
            detected_at.replace(tzinfo=timezone.utc).isoformat(timespec="seconds")
            if isinstance(detected_at, datetime)
            else str(detected_at or "")
        ),
        "status": row.get("status") or "-",
        "attackerIp": row.get("attacker_ip") or None,
        "requestUrl": row.get("request_url") or None,
        "ruleName": row.get("rule_name") or None,
        "blocked": (
            bool(row.get("blocked")) if row.get("blocked") is not None else None
        ),
        "blockResult": row.get("block_result") or None,
        "recommendation": row.get("recommendation") or None,
        "autoRemediation": bool(row.get("auto_remediation")),
        "highlightAssets": _json_list(row.get("highlight_assets")),
        "attackPath": _json_list(row.get("attack_path")),
        "logs": row.get("logs") or None,
    }


def _parse_query_datetime(value, label):
    try:
        naive = datetime.fromisoformat(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} 날짜/시간 형식이 올바르지 않습니다.") from exc
    # 프론트의 시작/종료 시간 입력(datetime-local)은 사용자가 보는 KST 기준값인데,
    # DB의 detected_at은 UTC라서 그대로 비교하면 9시간 어긋난다. KST로 해석해
    # UTC로 변환한 뒤 비교한다.
    return naive.replace(tzinfo=KST).astimezone(timezone.utc).replace(tzinfo=None)


def _log_time_window(args):
    range_key = str(args.get("range") or "").strip().lower()
    start_text = str(args.get("start") or "").strip()
    end_text = str(args.get("end") or "").strip()

    if range_key:
        if range_key not in LOG_RANGE_DELTAS:
            raise ValueError("지원하지 않는 조회 기간입니다.")
        end_at = datetime.now()
        return end_at - LOG_RANGE_DELTAS[range_key], end_at

    if not start_text or not end_text:
        raise ValueError("조회 기간 또는 시작/종료 시간을 지정해 주세요.")

    start_at = _parse_query_datetime(start_text, "시작")
    end_at = _parse_query_datetime(end_text, "종료")
    if start_at >= end_at:
        raise ValueError("시작 시간은 종료 시간보다 빨라야 합니다.")
    return start_at, end_at


def _read_security_logs(args):
    source = str(args.get("source") or "").strip().lower()
    if source not in LOG_SOURCES:
        raise ValueError("WAF, GuardDuty, Inspector 중 하나를 선택해 주세요.")

    start_at, end_at = _log_time_window(args)
    conditions = ["detected_at >= %s", "detected_at <= %s"]
    params = [start_at, end_at]

    normalized_scenario = "LOWER(COALESCE(scenario_type, ''))"
    if source == "waf":
        conditions.append(f"{normalized_scenario} IN (%s, %s, %s, %s)")
        params.extend(["sqli", "xss", "dir", "brute"])
    elif source == "guardduty":
        conditions.append(f"{normalized_scenario} IN (%s, %s)")
        params.extend(["port", "cred"])
    else:
        conditions.append(f"{normalized_scenario} = %s")
        params.append("vuln")

    severity = str(args.get("severity") or "").strip().capitalize()
    if severity:
        if severity not in VALID_SEVERITIES:
            raise ValueError("지원하지 않는 Severity입니다.")
        conditions.append("severity = %s")
        params.append(severity)

    scenario_type = str(
        args.get("scenario_type") or args.get("finding_type") or ""
    ).strip().lower()
    scenario_filters = {
        "sqli": (f"{normalized_scenario} = %s", ["sqli"]),
        "xss": (f"{normalized_scenario} = %s", ["xss"]),
        "dir": (f"{normalized_scenario} = %s", ["dir"]),
        "brute": (f"{normalized_scenario} = %s", ["brute"]),
        "port": (f"{normalized_scenario} = %s", ["port"]),
        "cred": (f"{normalized_scenario} = %s", ["cred"]),
    }
    if scenario_type:
        allowed_for_source = {
            "waf": {"sqli", "xss", "dir", "brute"},
            "guardduty": {"port", "cred"},
            "inspector": set(),
        }
        if scenario_type not in allowed_for_source[source]:
            raise ValueError("선택한 로그 종류에서 지원하지 않는 공격 유형입니다.")
        scenario_sql, scenario_params = scenario_filters[scenario_type]
        conditions.append(scenario_sql)
        params.extend(scenario_params)

    attacker_ip = str(args.get("attacker_ip") or "").strip()
    if attacker_ip:
        conditions.append("attacker_ip = %s")
        params.append(attacker_ip)

    asset = str(args.get("asset") or "").strip()
    if asset:
        conditions.append("asset LIKE %s")
        params.append(f"%{asset}%")

    action = str(args.get("action") or "").strip().upper()
    if action:
        if source != "waf" or action not in {"BLOCK", "ALLOW"}:
            raise ValueError("지원하지 않는 Action입니다.")
        conditions.append("blocked = %s")
        params.append(action == "BLOCK")

    cve = str(args.get("cve") or "").strip()
    if cve:
        if source != "inspector":
            raise ValueError("CVE 필터는 Inspector 로그에서만 사용할 수 있습니다.")
        cve_like = f"%{cve}%"
        conditions.append("(rule_name LIKE %s OR title LIKE %s OR logs LIKE %s)")
        params.extend([cve_like, cve_like, cve_like])

    query = f"""
        SELECT
            id, service, scenario_type, severity, title, asset, detected_at,
            status, attacker_ip, request_url, rule_name, blocked,
            block_result, recommendation, auto_remediation,
            highlight_assets, attack_path, logs
        FROM security_events
        WHERE {' AND '.join(conditions)}
        ORDER BY detected_at DESC
    """

    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(query, params)
            rows = cursor.fetchall()

    return [_event_to_log(row) for row in rows]


def _read_dashboard_data():
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT * FROM (
                    SELECT *,
                        ROW_NUMBER() OVER (
                            PARTITION BY scenario_type, COALESCE(attacker_ip, title)
                            ORDER BY detected_at DESC
                        ) AS rn,
                        COUNT(*) OVER (
                            PARTITION BY scenario_type, COALESCE(attacker_ip, title)
                        ) AS occurrence_count
                    FROM security_events
                    WHERE status NOT IN ('조치 완료', '자동 완료', '예외 처리', '완료')
                      AND scenario_type IN ('sqli', 'dir', 'brute', 'cred', 'vuln', 'xss', 'port')
                ) grouped
                WHERE rn = 1
                ORDER BY detected_at DESC
                """
            )
            action_rows = cursor.fetchall()

            cursor.execute(
                """
                SELECT *
                FROM security_events
                WHERE scenario_type IN ('sqli', 'dir', 'brute', 'cred', 'vuln', 'xss', 'port')
                ORDER BY detected_at DESC
                LIMIT 200
                """
            )
            detect_rows = cursor.fetchall()

            cursor.execute(
                """
                SELECT * FROM (
                    SELECT
                        rh.*,
                        se.title AS event_title,
                        se.asset AS asset,
                        se.attacker_ip AS attacker_ip,
                        se.scenario_type AS scenario_type,
                        ROW_NUMBER() OVER (
                            PARTITION BY se.scenario_type, COALESCE(se.attacker_ip, se.title)
                            ORDER BY COALESCE(rh.completed_at, rh.requested_at) DESC
                        ) AS rn,
                        COUNT(*) OVER (
                            PARTITION BY se.scenario_type, COALESCE(se.attacker_ip, se.title)
                        ) AS occurrence_count
                    FROM remediation_history rh
                    LEFT JOIN security_events se ON se.id = rh.event_id
                    WHERE se.scenario_type IN ('sqli', 'dir', 'brute', 'cred', 'vuln', 'xss', 'port')
                ) grouped
                WHERE rn = 1
                ORDER BY COALESCE(completed_at, requested_at) DESC
                LIMIT 200
                """
            )
            remediation_rows = cursor.fetchall()

    return {
        "events": [_event_to_action_event(row) for row in action_rows],
        "detectHistory": [_event_to_detect_history(row) for row in detect_rows],
        "remediationHistory": [
            _remediation_to_history(row) for row in remediation_rows
        ],
    }


def _iso_datetime(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc).isoformat().replace("+00:00", "Z")
    return str(value) if value else None


def _aggregate_service_metrics(rows):
    """service_metrics 는 5분 구간마다 서버별로 한 행씩(최대 5행) 쌓인다.
    같은 window_start 를 모아 시스템 전체 기준 한 점으로 합친다.
    - cpu/memory: 값이 있는 서버들의 평균
    - latency/rps/errorRate/healthy/unhealthy: ALB 뒤에 있는 서버(k3s, dashboard)만 더한다
      (shop_app/shop_db/security_db 는 ALB에 안 물려 있어 요청 지표가 없다)
    """
    by_window = {}
    for row in rows:
        by_window.setdefault(row["window_start"], []).append(row)

    points = []
    for window_start in sorted(by_window):
        server_rows = by_window[window_start]
        window_end = max(r["window_end"] for r in server_rows)

        cpu_values = [float(r["cpu_percent"]) for r in server_rows if r.get("cpu_percent") is not None]
        mem_values = [float(r["memory_percent"]) for r in server_rows if r.get("memory_percent") is not None]

        alb_rows = [r for r in server_rows if r.get("request_count") is not None]
        total_requests = sum(int(r["request_count"]) for r in alb_rows)

        lat_pairs = [
            (float(r["avg_latency_ms"]), int(r["request_count"]))
            for r in alb_rows
            if r.get("avg_latency_ms") is not None
        ]
        weight_sum = sum(w for _, w in lat_pairs)
        if lat_pairs and weight_sum > 0:
            latency = sum(v * w for v, w in lat_pairs) / weight_sum
        elif lat_pairs:
            latency = sum(v for v, _ in lat_pairs) / len(lat_pairs)
        else:
            latency = None

        if alb_rows:
            errors = sum(
                float(r["error_rate_percent"] or 0) / 100 * int(r["request_count"])
                for r in alb_rows
            )
            error_rate = (errors / total_requests * 100) if total_requests else 0.0
            rps = total_requests / 300  # 5분(300초) 구간 합계를 초당 요청 수로 환산
        else:
            error_rate = None
            rps = None

        healthy_vals = [int(r["healthy_targets"]) for r in alb_rows if r.get("healthy_targets") is not None]
        unhealthy_vals = [int(r["unhealthy_targets"]) for r in alb_rows if r.get("unhealthy_targets") is not None]
        healthy = sum(healthy_vals) if healthy_vals else None
        unhealthy = sum(unhealthy_vals) if unhealthy_vals else None

        statuses = {r["status"] for r in server_rows}
        if "unhealthy" in statuses:
            health = "CRITICAL"
        elif "degraded" in statuses:
            health = "WARNING"
        elif statuses and statuses <= {"healthy"}:
            health = "NORMAL"
        else:
            health = None

        points.append(
            {
                "window_start": window_start,
                "window_end": window_end,
                "cpu": sum(cpu_values) / len(cpu_values) if cpu_values else None,
                "memory": sum(mem_values) / len(mem_values) if mem_values else None,
                "latency": latency,
                "rps": rps,
                "error_rate": error_rate,
                "health": health,
                "healthy": healthy,
                "unhealthy": unhealthy,
            }
        )
    return points


def _read_overview_metrics():
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT server, cpu_percent, memory_percent, request_count,
                       avg_latency_ms, error_rate_percent, status,
                       healthy_targets, unhealthy_targets, window_start, window_end
                FROM service_metrics
                WHERE window_end >= UTC_TIMESTAMP() - INTERVAL 24 HOUR
                ORDER BY window_start ASC
                """
            )
            rows = cursor.fetchall()

    points = _aggregate_service_metrics(rows)[-OVERVIEW_SERIES_LIMIT:]

    def series(field):
        values = [p[field] for p in points if p[field] is not None]
        return {
            "current": round(values[-1], 3) if values else None,
            "series": [round(v, 3) for v in values],
            "collectedAt": _iso_datetime(points[-1]["window_end"]) if points else None,
        }

    cpu = series("cpu")
    memory = series("memory")
    latency = series("latency")
    rps = series("rps")
    error_rate = series("error_rate")

    health_points = [p for p in points if p["health"] is not None]
    last_health = health_points[-1] if health_points else None
    health = {
        "status": last_health["health"] if last_health else None,
        "healthy": last_health["healthy"] if last_health else None,
        "unhealthy": last_health["unhealthy"] if last_health else None,
        "collectedAt": _iso_datetime(last_health["window_end"]) if last_health else None,
    }

    return {
        "cpu": cpu,
        "memory": memory,
        "latency": latency,
        "rps": rps,
        "errorRate": error_rate,
        "health": health,
    }


def _monitoring_summary(records, field, unit):
    values = [record[field] for record in records if record.get(field) is not None]
    return {
        "current": round(values[-1], 3) if values else None,
        "average": round(sum(values) / len(values), 3) if values else None,
        "max": round(max(values), 3) if values else None,
        "min": round(min(values), 3) if values else None,
        "unit": unit,
    }




def _read_monitoring_metrics(args):
    requested = str(args.get("metric") or "all").strip().lower()
    metric = MONITORING_METRIC_ALIASES.get(requested)
    if metric is None:
        raise ValueError("지원하지 않는 운영 지표입니다.")

    start_at, end_at = _log_time_window(args)

    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT server, cpu_percent, memory_percent, request_count,
                       avg_latency_ms, error_rate_percent, status,
                       healthy_targets, unhealthy_targets, window_start, window_end
                FROM service_metrics
                WHERE window_end >= %s AND window_end <= %s
                ORDER BY window_start ASC
                """,
                (start_at, end_at),
            )
            rows = cursor.fetchall()

    # metric(all/cpu/memory/...) 선택은 아래에서 반환할 필드를 고르는 데만 쓰고,
    # 조회 자체는 항상 같은 표에서 전체 지표를 한 번에 가져온다.
    records = [
        {
            "timestamp": _iso_datetime(point["window_end"]),
            "cpu": point["cpu"],
            "memory": point["memory"],
            "latency": point["latency"],
            "rps": point["rps"],
            "errorRate": point["error_rate"],
            "health": point["health"],
            "healthy": point["healthy"],
            "unhealthy": point["unhealthy"],
        }
        for point in _aggregate_service_metrics(rows)
    ]

    units = {
        "cpu": "%",
        "memory": "%",
        "latency": "ms",
        "rps": "rps",
        "errorRate": "%",
    }
    summaries = {
        name: _monitoring_summary(records, name, unit)
        for name, unit in units.items()
    }
    health_records = [record for record in records if record["health"] is not None]
    health_summary = {
        "current": health_records[-1]["health"] if health_records else None,
        "healthy": health_records[-1]["healthy"] if health_records else None,
        "unhealthy": health_records[-1]["unhealthy"] if health_records else None,
    }

    if metric == "all":
        visible_records = [
            record
            for record in records
            if any(
                record[field] is not None
                for field in ("cpu", "memory", "latency", "rps", "errorRate", "health")
            )
        ]
        return {
            "metric": "all",
            "window": {"start": _iso_datetime(start_at), "end": _iso_datetime(end_at)},
            "summaries": {**summaries, "health": health_summary},
            "series": visible_records,
        }

    if metric == "health":
        return {
            "metric": "health",
            "window": {"start": _iso_datetime(start_at), "end": _iso_datetime(end_at)},
            "summary": health_summary,
            "series": health_records,
        }

    metric_records = [record for record in records if record[metric] is not None]
    return {
        "metric": metric,
        "window": {"start": _iso_datetime(start_at), "end": _iso_datetime(end_at)},
        "summary": summaries[metric],
        "series": [
            {"timestamp": record["timestamp"], "value": round(record[metric], 3)}
            for record in metric_records
        ],
    }


AI_DIAGNOSIS_ACTIVE_STATUSES = {"collecting", "diagnosing"}


def _ai_diagnosis_row_to_json(row):
    result = row.get("result")
    return {
        "id": row["id"],
        "status": row.get("status"),
        "message": row.get("message"),
        "requestedBy": row.get("requested_by"),
        "result": json.loads(result) if result else None,
        "error": row.get("error"),
        "startedAt": _format_datetime(row.get("started_at"), "%Y.%m.%d %H:%M:%S"),
        "finishedAt": _format_datetime(row.get("finished_at"), "%Y.%m.%d %H:%M:%S"),
    }


_AI_DIAGNOSIS_STATUS_FILL = {
    "PASS": PatternFill("solid", fgColor="C6E0B4"),
    "FAIL": PatternFill("solid", fgColor="F8CBAD"),
    "REVIEW": PatternFill("solid", fgColor="FFE699"),
    "N/A": PatternFill("solid", fgColor="D9D9D9"),
}
_AI_DIAGNOSIS_HEADER_FILL = PatternFill("solid", fgColor="305496")
_AI_DIAGNOSIS_HEADER_FONT = Font(color="FFFFFF", bold=True)
_AI_DIAGNOSIS_DETAIL_HEADERS = [
    "항목", "분류", "이름", "위험도", "판정",
    "현재 상태", "기대 상태", "판단 근거", "권장 조치",
]
_AI_DIAGNOSIS_DETAIL_WIDTHS = [8, 16, 30, 8, 10, 34, 28, 44, 34]


def _ai_diagnosis_style_header(ws, headers):
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = _AI_DIAGNOSIS_HEADER_FILL
        cell.font = _AI_DIAGNOSIS_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _ai_diagnosis_write_detail_rows(ws, results, rule_meta):
    for row in results:
        meta = rule_meta.get(str(row.get("rule_id")), {})
        ws.append([
            row.get("rule_id"),
            meta.get("category", ""),
            meta.get("name", ""),
            row.get("severity"),
            row.get("status"),
            row.get("current_value"),
            row.get("expected_value"),
            row.get("reason"),
            row.get("recommendation"),
        ])
        fill = _AI_DIAGNOSIS_STATUS_FILL.get(row.get("status"))
        if fill:
            ws.cell(row=ws.max_row, column=5).fill = fill
        ws.cell(row=ws.max_row, column=8).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=ws.max_row, column=9).alignment = Alignment(wrap_text=True, vertical="top")


def _build_ai_diagnosis_workbook(report):
    rule_meta = load_rule_meta()
    results = report.get("results") or []
    summary = report.get("summary") or {}

    wb = Workbook()

    ws = wb.active
    ws.title = "요약"
    ws.append(["AWS 보안 구성 AI 진단 결과"])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([f"기준: {report.get('standard') or '-'}"])
    ws.append([f"진단 모델: {report.get('model') or '-'}"])
    ws.append([f"수집 시각(UTC): {report.get('collected_at') or '-'}"])
    ws.append([f"리전: {report.get('region') or '-'}"])
    ws.append([])

    ws.append(["구분", "건수"])
    _ai_diagnosis_style_header(ws, ["구분", "건수"])
    for label, key in [
        ("PASS", "pass"), ("FAIL", "fail"), ("REVIEW", "review"),
        ("N/A", "na"), ("전체", "total"),
    ]:
        ws.append([label, summary.get(key, 0)])
        fill = _AI_DIAGNOSIS_STATUS_FILL.get(label)
        if fill:
            ws.cell(row=ws.max_row, column=1).fill = fill

    ws.append([])
    ws.append(["AI 종합 소견"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([report.get("consultant_comment") or "(생성되지 않음)"])
    ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=6)
    ws.row_dimensions[ws.max_row].height = 90

    for col, width in enumerate([20, 60, 16, 16, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws_detail = wb.create_sheet("상세 결과")
    ws_detail.append(_AI_DIAGNOSIS_DETAIL_HEADERS)
    _ai_diagnosis_style_header(ws_detail, _AI_DIAGNOSIS_DETAIL_HEADERS)
    _ai_diagnosis_write_detail_rows(ws_detail, results, rule_meta)
    for col, width in enumerate(_AI_DIAGNOSIS_DETAIL_WIDTHS, start=1):
        ws_detail.column_dimensions[get_column_letter(col)].width = width

    ws_action = wb.create_sheet("조치 필요")
    ws_action.append(_AI_DIAGNOSIS_DETAIL_HEADERS)
    _ai_diagnosis_style_header(ws_action, _AI_DIAGNOSIS_DETAIL_HEADERS)
    _ai_diagnosis_write_detail_rows(
        ws_action,
        [r for r in results if r.get("status") in ("FAIL", "REVIEW")],
        rule_meta,
    )
    for col, width in enumerate(_AI_DIAGNOSIS_DETAIL_WIDTHS, start=1):
        ws_action.column_dimensions[get_column_letter(col)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# AI 진단 PDF 보고서. reportlab 기본 내장 폰트엔 한글 글리프가 없어서, 실제
# 폰트 파일(나눔고딕, Dockerfile에서 apt로 설치)을 등록해서 써야 한다.
# 폰트 파일이 없는 환경(예: 로컬 개발)에서는 기본 폰트로 대체하되, 한글은
# 깨질 수 있다 - PDF 생성 자체가 실패하지는 않게만 해둔다.
# ---------------------------------------------------------------------------
_KOREAN_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    "/usr/share/fonts/opentype/nanum/NanumGothic.ttf",
]
_KOREAN_FONT_BOLD_CANDIDATES = [
    "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf",
    "/usr/share/fonts/opentype/nanum/NanumGothicBold.ttf",
]


def _register_korean_pdf_font():
    regular_path = next((p for p in _KOREAN_FONT_CANDIDATES if os.path.exists(p)), None)
    if not regular_path:
        app.logger.warning("한글 PDF 폰트를 찾지 못했습니다 - 기본 폰트로 대체됩니다.")
        return "Helvetica", "Helvetica-Bold"

    pdfmetrics.registerFont(TTFont("NanumGothic", regular_path))
    bold_path = next((p for p in _KOREAN_FONT_BOLD_CANDIDATES if os.path.exists(p)), None)
    if bold_path:
        pdfmetrics.registerFont(TTFont("NanumGothic-Bold", bold_path))
        return "NanumGothic", "NanumGothic-Bold"
    return "NanumGothic", "NanumGothic"


_PDF_FONT, _PDF_FONT_BOLD = _register_korean_pdf_font()

_PDF_STATUS_HEX = {
    "PASS": "#067647", "FAIL": "#B42318", "REVIEW": "#B54708", "N/A": "#667085",
}
_PDF_STATUS_BG = {
    "PASS": colors.HexColor("#ECFDF3"),
    "FAIL": colors.HexColor("#FEF3F2"),
    "REVIEW": colors.HexColor("#FFFAEB"),
    "N/A": colors.HexColor("#F2F4F7"),
}

_pdf_style_title = ParagraphStyle(
    "AIDiagTitle", fontName=_PDF_FONT_BOLD, fontSize=18, leading=22,
)
_pdf_style_meta = ParagraphStyle(
    "AIDiagMeta", fontName=_PDF_FONT, fontSize=9, leading=13,
    textColor=colors.HexColor("#667085"),
)
_pdf_style_h2 = ParagraphStyle(
    "AIDiagH2", fontName=_PDF_FONT_BOLD, fontSize=13, leading=17,
    spaceBefore=14, spaceAfter=6, textColor=colors.HexColor("#101828"),
)
_pdf_style_body = ParagraphStyle(
    "AIDiagBody", fontName=_PDF_FONT, fontSize=9.5, leading=14,
    textColor=colors.HexColor("#344054"),
)
_pdf_style_kv_title = ParagraphStyle(
    "AIDiagKVTitle", fontName=_PDF_FONT_BOLD, fontSize=10.5, leading=14,
    textColor=colors.HexColor("#101828"),
)
_pdf_style_kv_value = ParagraphStyle(
    "AIDiagKVValue", fontName=_PDF_FONT, fontSize=9, leading=13,
    textColor=colors.HexColor("#344054"),
)


def _pdf_esc(value):
    return _xml_escape(str(value or ""))


def _build_ai_diagnosis_pdf(report):
    rule_meta = load_rule_meta()
    results = report.get("results") or []
    summary = report.get("summary") or {}

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
    )
    story = [
        Paragraph("AWS 보안 구성 AI 진단 결과", _pdf_style_title),
        Spacer(1, 4),
    ]
    for line in [
        f"기준: {_pdf_esc(report.get('standard') or '-')}",
        f"진단 모델: {_pdf_esc(report.get('model') or '-')}",
        f"수집 시각(UTC): {_pdf_esc(report.get('collected_at') or '-')}",
        f"리전: {_pdf_esc(report.get('region') or '-')}",
    ]:
        story.append(Paragraph(line, _pdf_style_meta))
    story.append(Spacer(1, 10))

    summary_table = Table(
        [
            ["PASS", "FAIL", "REVIEW", "N/A", "전체"],
            [
                str(summary.get("pass", 0)),
                str(summary.get("fail", 0)),
                str(summary.get("review", 0)),
                str(summary.get("na", 0)),
                str(summary.get("total", 0)),
            ],
        ],
        colWidths=[30 * mm] * 5,
    )
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#305496")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, -1), _PDF_FONT),
                ("FONTNAME", (0, 0), (-1, 0), _PDF_FONT_BOLD),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D5DD")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("BACKGROUND", (0, 1), (0, 1), _PDF_STATUS_BG["PASS"]),
                ("BACKGROUND", (1, 1), (1, 1), _PDF_STATUS_BG["FAIL"]),
                ("BACKGROUND", (2, 1), (2, 1), _PDF_STATUS_BG["REVIEW"]),
                ("BACKGROUND", (3, 1), (3, 1), _PDF_STATUS_BG["N/A"]),
            ]
        )
    )
    story.append(summary_table)
    story.append(Spacer(1, 12))

    if report.get("consultant_comment"):
        story.append(Paragraph("AI 종합 소견", _pdf_style_h2))
        story.append(Paragraph(_pdf_esc(report["consultant_comment"]), _pdf_style_body))

    by_category = {}
    for row in results:
        meta = rule_meta.get(str(row.get("rule_id")), {})
        by_category.setdefault(meta.get("category") or "기타", []).append((row, meta))

    for category in CATEGORY_ORDER:
        rows = by_category.get(category)
        if not rows:
            continue
        story.append(Paragraph(_pdf_esc(category), _pdf_style_h2))
        for row, meta in rows:
            status = row.get("status")
            status_hex = _PDF_STATUS_HEX.get(status, "#101828")

            title_para = Paragraph(
                f"{_pdf_esc(row.get('rule_id'))} {_pdf_esc(meta.get('name'))}",
                _pdf_style_kv_title,
            )
            status_para = Paragraph(
                f"<font color='{status_hex}'><b>{_pdf_esc(status)}</b></font>"
                f"&nbsp;&nbsp;(위험도 {_pdf_esc(row.get('severity'))})",
                _pdf_style_kv_value,
            )

            item_rows = [["항목", title_para], ["판정", status_para]]
            for label, key in [
                ("현재 상태", "current_value"),
                ("기대 상태", "expected_value"),
                ("판단 근거", "reason"),
                ("권장 조치", "recommendation"),
            ]:
                value = row.get(key)
                if value:
                    item_rows.append(
                        [label, Paragraph(_pdf_esc(value), _pdf_style_kv_value)]
                    )

            item_table = Table(item_rows, colWidths=[26 * mm, 148 * mm])
            item_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F9FAFB")),
                        ("FONTNAME", (0, 0), (0, -1), _PDF_FONT_BOLD),
                        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#667085")),
                        ("FONTSIZE", (0, 0), (0, -1), 8.5),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E4E7EC")),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        (
                            "BACKGROUND",
                            (1, 1),
                            (1, 1),
                            _PDF_STATUS_BG.get(status, colors.white),
                        ),
                    ]
                )
            )
            story.append(KeepTogether([item_table, Spacer(1, 8)]))

    doc.build(story)
    return buffer.getvalue()


def _run_ai_diagnosis_job(run_id):
    """collect_aws_state()/diagnose_aws_state()는 IAM 전체 조회 + OpenAI 호출
    4번을 순서대로 돌기 때문에 수십 초~수 분이 걸릴 수 있다. 요청 스레드를
    붙잡지 않도록 백그라운드 스레드에서 실행하고, 진행 상황은 gunicorn
    워커가 여러 개라도 모두 같은 값을 보도록 DB(ai_diagnosis_runs)에 남긴다."""

    def _update(**fields):
        columns = ", ".join(f"{key} = %s" for key in fields)
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    f"UPDATE ai_diagnosis_runs SET {columns} WHERE id = %s",
                    (*fields.values(), run_id),
                )

    try:
        _update(status="collecting", message="AWS 현재 설정을 수집하는 중입니다...")
        aws_state = collect_aws_state()

        _update(status="diagnosing", message="AI가 33개 항목을 진단하는 중입니다...")
        result = diagnose_aws_state(aws_state)

        _update(
            status="done",
            message=None,
            result=json.dumps(result, ensure_ascii=False, default=str),
            finished_at=datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0),
        )
    except DiagnosisError as exc:
        app.logger.exception("AI 진단 실패(run_id=%s)", run_id)
        _update(
            status="error",
            message=None,
            error=str(exc),
            finished_at=datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0),
        )
    except Exception as exc:
        app.logger.exception("AI 진단 실패, 예상치 못한 오류(run_id=%s)", run_id)
        _update(
            status="error",
            message=None,
            error=f"예상치 못한 오류: {exc}",
            finished_at=datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0),
        )


@app.get("/api/health")
def health():
    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute("SELECT 1 AS ok")
                result = cursor.fetchone()
        return jsonify({"status": "ok", "db": result["ok"] == 1})
    except Exception as exc:
        app.logger.exception("DB health check failed")
        return jsonify({"status": "error", "message": str(exc)}), 500


@app.get("/api/auth/status")
def auth_status():
    if not session.get("authenticated"):
        return jsonify({"authenticated": False})
    return jsonify({"authenticated": True, "user": _current_user()})


def _login_accounts():
    """관리자 계정은 항상 존재하고, 승인자 계정은 APPROVER_PASSWORD가
    설정된 경우에만 로그인 가능한 계정으로 취급한다(둘 다 env var 기반
    단일 공유 계정 - 여러 명의 승인자를 개별 관리하는 건 아직 아님)."""
    accounts = [
        {
            "username": os.getenv("ADMIN_USERNAME", "admin"),
            "password": os.getenv("ADMIN_PASSWORD", ""),
            "role": os.getenv("ADMIN_ROLE", "관리자"),
            "team": os.getenv("ADMIN_TEAM", "보안관제팀"),
        }
    ]
    if os.getenv("APPROVER_PASSWORD", "").strip():
        accounts.append(
            {
                "username": os.getenv("APPROVER_USERNAME", "approver"),
                "password": os.getenv("APPROVER_PASSWORD", ""),
                "role": os.getenv("APPROVER_ROLE", "승인자"),
                "team": os.getenv("APPROVER_TEAM", "보안관제팀"),
            }
        )
    return accounts


@app.post("/api/auth/login")
def auth_login():
    data = request.get_json(silent=True) or {}
    username = str(data.get("username") or "").strip()
    password = str(data.get("password") or "")

    accounts = _login_accounts()
    if not any(acc["password"] for acc in accounts):
        return (
            jsonify(
                {
                    "error": "AUTH_NOT_CONFIGURED",
                    "message": "backend/.env에 ADMIN_PASSWORD를 먼저 설정해 주세요.",
                }
            ),
            503,
        )

    matched = None
    for account in accounts:
        if not account["password"]:
            continue
        username_ok = hmac.compare_digest(username, account["username"])
        password_ok = hmac.compare_digest(password, account["password"])
        if username_ok and password_ok:
            matched = account
            break

    if not matched:
        return jsonify({"error": "INVALID_CREDENTIALS", "message": "아이디 또는 비밀번호가 올바르지 않습니다."}), 401

    session.clear()
    session.permanent = True
    session["authenticated"] = True
    session["username"] = matched["username"]
    session["role"] = matched["role"]
    session["team"] = matched["team"]

    return jsonify({"ok": True, "user": _current_user()})


@app.post("/api/auth/logout")
def auth_logout():
    session.clear()
    return jsonify({"ok": True})


APPROVAL_PENDING = "대기"
APPROVAL_APPROVED = "승인"
APPROVAL_REJECTED = "반려"
# 주의: Lambda A(mapping.py)가 자동 조치 가능한 탐지 건에 기본으로 붙이는
# 상태가 이미 "승인 대기"다(관리자가 아직 안 눌렀다는 뜻, infra 레포에서
# 정의됨). 여기서 같은 문자열을 쓰면 "그냥 아직 아무도 안 건드린 탐지 건"과
# "승인자한테 요청을 실제로 보낸 건"을 구분할 수 없게 된다 - 실제로 이 충돌
# 때문에 요청을 하나도 안 보낸 자동조치 대상 건들까지 전부 승인 대기 중인
# 것처럼 잠겨버리는 문제가 있었다. 겹치지 않는 별도 문자열을 쓴다.
PENDING_APPROVAL_STATUS = "승인 요청됨"
# 수동 조치가 승인자 승인까지는 끝났지만, 실제로 사람이 손으로 그 조치를
# 아직 안 한 상태. 관리자가 "수동 조치 완료 처리"를 눌러야 비로소
# '조치 완료'로 넘어간다(자동 조치는 승인 즉시 Lambda가 실행하니 이 중간
# 단계가 없다).
MANUAL_APPROVED_STATUS = "수동 조치 대기"


def _execute_remediation_action(event, action, approver_username):
    """실제 조치(block_ip/disable_access_key 등)를 Lambda로 실행하고, 같은
    그룹으로 화면에 합쳐서 보이던 나머지 중복 이벤트도 같이 닫는다.
    성공 여부와 실패 메시지를 반환한다(예외를 던지지 않음 - 호출하는 쪽이
    승인 요청 상태를 같이 다루기 때문에 흐름 제어를 명시적으로 하려는 것)."""
    event_id = event["id"]
    function_name = os.getenv("REMEDIATION_LAMBDA_NAME", "").strip()
    if not function_name:
        return False, "REMEDIATION_LAMBDA_NAME 설정이 필요합니다."

    payload = {
        "event_id": event_id, "action": action,
        "approver": approver_username, "method": "수동", "params": {},
    }
    try:
        # Mutating invocation must not be retried automatically after an ambiguous timeout.
        client = boto3.client(
            "lambda", region_name=os.getenv("AWS_REGION", "ap-northeast-2"),
            config=Config(read_timeout=900, connect_timeout=10, retries={"total_max_attempts": 1}),
        )
        response = client.invoke(FunctionName=function_name, InvocationType="RequestResponse",
                                 Payload=json.dumps(payload, ensure_ascii=False).encode("utf-8"))
        result = json.loads(response["Payload"].read())
        if not isinstance(result, dict):
            raise ValueError("Lambda 응답이 JSON 객체가 아닙니다.")
        body = result.get("body", result)
        if isinstance(body, str):
            body = json.loads(body)
        if not isinstance(body, dict):
            raise ValueError("Lambda 응답 body가 JSON 객체가 아닙니다.")
        status = result.get("statusCode")
        success = ((type(status) is int and 200 <= status < 300)
                   or (status is None and body.get("success") is True))
        if (response.get("FunctionError") or response.get("StatusCode") != 200
                or not success or result.get("success") is False or result.get("error")
                or body.get("success") is False or body.get("error")
                or str(body.get("status", "")).lower() in ("failed", "failure", "error", "실패")
                or str(body.get("result", "")).lower() in ("failed", "failure", "error", "실패")):
            message = body.get("message") or body.get("errorMessage") or body.get("error") or "Lambda 조치 실행에 실패했습니다."
            return False, str(message)
    except Exception as exc:
        app.logger.exception("Remediation invocation failed")
        return False, str(exc)

    # Lambda alone owns event_id 자신의 status/remediation_history 기록.
    # 화면에 이 이벤트 하나로 합쳐서 보이던 나머지(같은 IP/같은 취약점의 반복 탐지)도
    # 같은 조치가 이미 적용된 셈이니 같이 닫는다 - 실패해도 대표 건 자체는 이미
    # 성공했으므로 전체 요청을 실패시키지 않는다(다음 새로고침 때 남은 게 있으면
    # 그것만 다시 보일 뿐).
    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                group_ids = [
                    gid for gid in _open_group_ids(
                        cursor, event.get("scenario_type"), event.get("attacker_ip"), event.get("title")
                    ) if gid != event_id
                ]
                for group_id in group_ids:
                    cursor.execute(
                        "UPDATE security_events SET status='조치 완료'"
                        + (", blocked=TRUE, block_result='성공'" if action == "block_ip" else "")
                        + " WHERE id = %s", (group_id,)
                    )
                    cursor.execute(
                        "INSERT INTO remediation_history "
                        "(event_id, action_type, method, approver, status, result, completed_at) "
                        "VALUES (%s, %s, '자동', %s, '완료', '동일 조치로 일괄 처리', UTC_TIMESTAMP())",
                        (group_id, action, approver_username),
                    )
    except Exception:
        app.logger.exception("Failed to close grouped duplicate events after remediation")

    return True, None


def _approval_request_to_json(row):
    return {
        "id": row["id"],
        "eventId": row.get("event_id"),
        "eventTitle": row.get("event_title"),
        "eventScenarioType": row.get("event_scenario_type"),
        "eventSeverity": row.get("event_severity"),
        "eventAsset": row.get("event_asset"),
        "eventStatus": row.get("event_status"),
        "actionType": row.get("action_type"),
        "note": row.get("note"),
        "requestType": row.get("request_type"),
        "status": row.get("status"),
        "requestedBy": row.get("requested_by"),
        "reviewedBy": row.get("reviewed_by"),
        "rejectReason": row.get("reject_reason"),
        "requestedAt": _format_datetime(row.get("requested_at"), "%Y.%m.%d %H:%M"),
        "reviewedAt": _format_datetime(row.get("reviewed_at"), "%Y.%m.%d %H:%M"),
        "needsManualCompletion": (
            row.get("request_type") == "manual"
            and row.get("status") == APPROVAL_APPROVED
            and row.get("event_status") == MANUAL_APPROVED_STATUS
        ),
    }


@app.post("/api/approval-requests")
@role_required(os.getenv("ADMIN_ROLE", "관리자"))
def create_approval_request():
    """관리자가 조치(수동/자동 모두)를 승인자에게 요청한다. 여기서는 아무
    조치도 실행하지 않는다 - 승인자가 승인해야만 그때 실제로 실행된다."""
    data = request.get_json(silent=True) or {}
    event_id = str(data.get("event_id") or "").strip()
    if not event_id:
        return jsonify(error="INVALID_REQUEST", message="event_id가 필요합니다."), 400
    note = str(data.get("note") or "").strip() or None

    requested_by = session.get("username") or "admin"

    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute("SELECT * FROM security_events WHERE id = %s", (event_id,))
                event = cursor.fetchone()
                if not event:
                    return jsonify(error="EVENT_NOT_FOUND", message="보안 이벤트를 찾을 수 없습니다."), 404
                if event.get("status") in REMEDIATION_CLOSED_STATUSES:
                    return jsonify(error="EVENT_ALREADY_CLOSED", message="이미 완료되거나 예외 처리된 이벤트입니다."), 409
                if event.get("status") == PENDING_APPROVAL_STATUS:
                    return jsonify(error="ALREADY_REQUESTED", message="이미 승인자에게 보낸 요청이 있습니다."), 409

                action = REMEDIATION_ACTIONS.get(event.get("scenario_type"))
                request_type = "auto" if action else "manual"
                if action == "block_ip" and not str(event.get("attacker_ip") or "").strip():
                    return jsonify(error="REMEDIATION_DATA_MISSING", message="차단할 공격 IP 정보가 없습니다."), 400

                cursor.execute(
                    "INSERT INTO approval_requests "
                    "(event_id, action_type, note, request_type, status, previous_status, requested_by, requested_at) "
                    "VALUES (%s, %s, %s, %s, %s, %s, %s, UTC_TIMESTAMP())",
                    (event_id, action, note, request_type, APPROVAL_PENDING, event.get("status"), requested_by),
                )
                request_id = cursor.lastrowid

                # 화면에 하나로 합쳐서 보이던 나머지(같은 종류로 반복 감지된 것들)도
                # 같이 PENDING_APPROVAL_STATUS로 표시한다 - 안 그러면 다음 새로고침 때
                # 그중 하나가 새 대표 건으로 튀어나와 마치 요청이 씹힌 것처럼 보인다.
                group_ids = _open_group_ids(
                    cursor, event.get("scenario_type"), event.get("attacker_ip"), event.get("title")
                )
                for group_id in group_ids:
                    cursor.execute(
                        "UPDATE security_events SET status = %s WHERE id = %s",
                        (PENDING_APPROVAL_STATUS, group_id),
                    )
    except Exception:
        app.logger.exception("Failed to create approval request")
        return jsonify(error="DATABASE_ERROR", message="승인 요청 생성 중 오류가 발생했습니다."), 500

    return jsonify(success=True, requestId=request_id)


@app.get("/api/approval-requests")
@login_required
def list_approval_requests():
    status_filter = str(request.args.get("status") or "").strip()
    query = (
        "SELECT ar.*, se.title AS event_title, se.scenario_type AS event_scenario_type, se.severity AS event_severity, "
        "se.asset AS event_asset, se.status AS event_status "
        "FROM approval_requests ar "
        "LEFT JOIN security_events se ON se.id = ar.event_id "
    )
    params = []
    if status_filter == APPROVAL_PENDING:
        # "대기 중" 화면은 "내가 지금 뭔가 해야 하는 것"을 보여주는 화면이라,
        # 아직 승인자 판단이 안 난 것(ar.status='대기')뿐 아니라 승인은
        # 됐지만 관리자가 아직 손으로 완료 처리를 안 한 수동 조치도 같이 보여준다.
        query += (
            "WHERE (ar.status = %s "
            "OR (ar.status = %s AND ar.request_type = 'manual' AND se.status = %s)) "
        )
        params.extend([APPROVAL_PENDING, APPROVAL_APPROVED, MANUAL_APPROVED_STATUS])
    elif status_filter:
        query += "WHERE ar.status = %s "
        params.append(status_filter)
    query += "ORDER BY ar.requested_at DESC LIMIT 200"

    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(query, params)
                rows = cursor.fetchall()
    except Exception:
        app.logger.exception("Failed to list approval requests")
        return jsonify(error="DATABASE_ERROR", message="승인 요청 목록을 조회하지 못했습니다."), 500

    return jsonify(requests=[_approval_request_to_json(row) for row in rows])


@app.post("/api/approval-requests/<int:request_id>/approve")
@role_required(os.getenv("APPROVER_ROLE", "승인자"))
def approve_approval_request(request_id):
    approver = session.get("username") or "approver"

    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute("SELECT * FROM approval_requests WHERE id = %s", (request_id,))
                req = cursor.fetchone()
    except Exception:
        app.logger.exception("Failed to read approval request")
        return jsonify(error="DATABASE_ERROR", message="승인 요청을 조회하지 못했습니다."), 500

    if not req:
        return jsonify(error="REQUEST_NOT_FOUND", message="승인 요청을 찾을 수 없습니다."), 404
    if req["status"] != APPROVAL_PENDING:
        return jsonify(error="REQUEST_ALREADY_REVIEWED", message="이미 처리된 요청입니다."), 409

    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute("SELECT * FROM security_events WHERE id = %s", (req["event_id"],))
                event = cursor.fetchone()
    except Exception:
        app.logger.exception("Failed to read event for approval")
        return jsonify(error="DATABASE_ERROR", message="보안 이벤트를 조회하지 못했습니다."), 500
    if not event:
        return jsonify(error="EVENT_NOT_FOUND", message="보안 이벤트를 찾을 수 없습니다."), 404

    if req["request_type"] == "auto":
        success, error_message = _execute_remediation_action(event, req.get("action_type"), approver)
        if not success:
            return jsonify(error="REMEDIATION_FAILED", message=error_message), 502
    else:
        # 수동 조치는 승인만으로 끝이 아니다 - 실제로 사람이 손으로 그 조치를
        # 해야 하므로, 여기서는 "이 계획대로 진행해도 좋다"는 승인만 기록하고
        # MANUAL_APPROVED_STATUS로 옮겨둔다. '조치 완료'로 넘어가는 건 관리자가
        # 실제로 조치를 마치고 /complete 를 호출할 때다.
        try:
            with get_connection() as connection:
                with connection.cursor() as cursor:
                    group_ids = _open_group_ids(
                        cursor, event.get("scenario_type"), event.get("attacker_ip"), event.get("title")
                    )
                    for group_id in group_ids:
                        cursor.execute(
                            "UPDATE security_events SET status = %s WHERE id = %s",
                            (MANUAL_APPROVED_STATUS, group_id),
                        )
        except Exception:
            app.logger.exception("Failed to move manual remediation event to approved state")
            return jsonify(error="DATABASE_ERROR", message="이벤트 상태 갱신 중 오류가 발생했습니다."), 500

    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    "UPDATE approval_requests SET status = %s, reviewed_by = %s, "
                    "reviewed_at = UTC_TIMESTAMP() WHERE id = %s",
                    (APPROVAL_APPROVED, approver, request_id),
                )
    except Exception:
        app.logger.exception("Failed to mark approval request as approved")

    return jsonify(success=True)


@app.post("/api/approval-requests/<int:request_id>/complete")
@role_required(os.getenv("ADMIN_ROLE", "관리자"))
def complete_manual_approval_request(request_id):
    """승인자가 승인한 수동 조치를, 실제로 손으로 다 마친 뒤 관리자가 눌러서
    '조치 완료'로 넘긴다. 자동 조치는 승인 시점에 Lambda가 바로 실행하니
    이 단계가 필요 없다(request_type이 auto면 거절)."""
    approver = session.get("username") or "admin"

    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute("SELECT * FROM approval_requests WHERE id = %s", (request_id,))
                req = cursor.fetchone()
    except Exception:
        app.logger.exception("Failed to read approval request for completion")
        return jsonify(error="DATABASE_ERROR", message="승인 요청을 조회하지 못했습니다."), 500

    if not req:
        return jsonify(error="REQUEST_NOT_FOUND", message="승인 요청을 찾을 수 없습니다."), 404
    if req["request_type"] != "manual" or req["status"] != APPROVAL_APPROVED:
        return jsonify(error="NOT_COMPLETABLE", message="완료 처리할 수 있는 상태가 아닙니다."), 409

    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute("SELECT * FROM security_events WHERE id = %s", (req["event_id"],))
                event = cursor.fetchone()
                if not event:
                    return jsonify(error="EVENT_NOT_FOUND", message="보안 이벤트를 찾을 수 없습니다."), 404
                if event.get("status") != MANUAL_APPROVED_STATUS:
                    return jsonify(error="NOT_COMPLETABLE", message="완료 처리할 수 있는 상태가 아닙니다."), 409

                group_ids = _open_group_ids(
                    cursor, event.get("scenario_type"), event.get("attacker_ip"), event.get("title")
                )
                for group_id in group_ids:
                    cursor.execute(
                        "UPDATE security_events SET status = '조치 완료' WHERE id = %s",
                        (group_id,),
                    )
                    cursor.execute(
                        "INSERT INTO remediation_history "
                        "(event_id, action_type, method, approver, status, result, completed_at) "
                        "VALUES (%s, %s, '수동', %s, '완료', %s, UTC_TIMESTAMP())",
                        (
                            group_id,
                            req.get("action_type") or "manual",
                            approver,
                            req.get("note") or "수동 조치 완료 처리",
                        ),
                    )
    except Exception:
        app.logger.exception("Failed to complete manual remediation")
        return jsonify(error="DATABASE_ERROR", message="완료 처리 중 오류가 발생했습니다."), 500

    return jsonify(success=True)


@app.post("/api/approval-requests/<int:request_id>/reject")
@role_required(os.getenv("APPROVER_ROLE", "승인자"))
def reject_approval_request(request_id):
    data = request.get_json(silent=True) or {}
    reason = str(data.get("reason") or "").strip() or None
    approver = session.get("username") or "approver"

    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute("SELECT * FROM approval_requests WHERE id = %s", (request_id,))
                req = cursor.fetchone()
                if not req:
                    return jsonify(error="REQUEST_NOT_FOUND", message="승인 요청을 찾을 수 없습니다."), 404
                if req["status"] != APPROVAL_PENDING:
                    return jsonify(error="REQUEST_ALREADY_REVIEWED", message="이미 처리된 요청입니다."), 409

                cursor.execute(
                    "SELECT scenario_type, attacker_ip, title FROM security_events WHERE id = %s",
                    (req["event_id"],),
                )
                event = cursor.fetchone()

                restore_status = req.get("previous_status") or "검토 필요"
                # 요청 당시 같이 PENDING_APPROVAL_STATUS로 묶었던 중복 건들도 원래
                # 상태로 되돌려서 관리자가 다시 조치 요청을 보낼 수 있게 한다("재승인요청").
                group_ids = (
                    _open_group_ids(cursor, event["scenario_type"], event.get("attacker_ip"), event.get("title"))
                    if event else [req["event_id"]]
                )
                for group_id in group_ids:
                    cursor.execute(
                        "UPDATE security_events SET status = %s WHERE id = %s",
                        (restore_status, group_id),
                    )

                cursor.execute(
                    "UPDATE approval_requests SET status = %s, reviewed_by = %s, "
                    "reviewed_at = UTC_TIMESTAMP(), reject_reason = %s WHERE id = %s",
                    (APPROVAL_REJECTED, approver, reason, request_id),
                )
    except Exception:
        app.logger.exception("Failed to reject approval request")
        return jsonify(error="DATABASE_ERROR", message="반려 처리 중 오류가 발생했습니다."), 500

    return jsonify(success=True)


@app.post("/api/exception")
@login_required
def except_events():
    """선택한 이벤트(1개 이상)를 '예외 처리' 상태로 바꾼다. 실제 조치(WAF 차단 등)는
    실행하지 않고, "이건 검토했고 조치 안 해도 된다"고 기록만 남기는 것."""
    data = request.get_json(silent=True)
    event_ids = (data or {}).get("event_ids")
    if (not isinstance(event_ids, list) or not event_ids
            or not all(isinstance(e, str) and e.strip() for e in event_ids)):
        return jsonify(error="INVALID_REQUEST", message="event_ids는 비어있지 않은 문자열 배열이어야 합니다."), 400
    if len(event_ids) > 200:
        return jsonify(error="INVALID_REQUEST", message="한 번에 최대 200건까지 처리할 수 있습니다."), 400

    approver = session.get("username") or "admin"
    updated, skipped = [], []
    try:
        with get_connection() as connection:
            with connection.cursor() as cursor:
                to_close = {}
                for event_id in event_ids:
                    cursor.execute(
                        "SELECT id, status, scenario_type, attacker_ip, title "
                        "FROM security_events WHERE id = %s", (event_id,)
                    )
                    event = cursor.fetchone()
                    if not event or event["status"] in REMEDIATION_CLOSED_STATUSES:
                        skipped.append(event_id)
                        continue
                    # 화면에 하나로 합쳐서 보이던 나머지(같은 종류로 반복 감지된 것들)도
                    # 같이 예외 처리한다.
                    for group_id in _open_group_ids(
                        cursor, event["scenario_type"], event.get("attacker_ip"), event.get("title")
                    ):
                        to_close[group_id] = True

                for group_id in to_close:
                    cursor.execute(
                        "UPDATE security_events SET status = '예외 처리' WHERE id = %s", (group_id,)
                    )
                    cursor.execute(
                        "INSERT INTO remediation_history "
                        "(event_id, action_type, method, approver, status, result, completed_at) "
                        "VALUES (%s, 'exception', '수동', %s, '완료', '예외 처리', UTC_TIMESTAMP())",
                        (group_id, approver),
                    )
                    updated.append(group_id)
    except Exception:
        app.logger.exception("Failed to mark events as exception")
        return jsonify(error="DATABASE_ERROR", message="예외 처리 중 오류가 발생했습니다."), 500

    return jsonify(success=True, updated=updated, skipped=skipped)


@app.get("/api/dashboard")
@login_required
def dashboard():
    try:
        return jsonify(_read_dashboard_data())
    except Exception as exc:
        app.logger.exception("Failed to read dashboard data")
        return jsonify({"error": "DATABASE_READ_FAILED", "message": str(exc)}), 500


@app.get("/api/events")
@login_required
def events():
    try:
        return jsonify(_read_dashboard_data()["events"])
    except Exception as exc:
        app.logger.exception("Failed to read events")
        return jsonify({"error": "DATABASE_READ_FAILED", "message": str(exc)}), 500


@app.get("/api/overview-metrics")
@login_required
def overview_metrics():
    try:
        return jsonify(_read_overview_metrics())
    except Exception as exc:
        app.logger.exception("Failed to read overview metrics")
        return jsonify({"error": "DATABASE_READ_FAILED", "message": str(exc)}), 500


@app.get("/api/logs")
@login_required
def security_logs():
    try:
        logs = _read_security_logs(request.args)
        return jsonify({"logs": logs, "count": len(logs)})
    except ValueError as exc:
        return jsonify({"error": "INVALID_QUERY", "message": str(exc)}), 400
    except Exception as exc:
        app.logger.exception("Failed to read security logs")
        return jsonify({"error": "DATABASE_READ_FAILED", "message": str(exc)}), 500


@app.get("/api/monitoring/metrics")
@login_required
def monitoring_metrics():
    try:
        return jsonify(_read_monitoring_metrics(request.args))
    except ValueError as exc:
        return jsonify({"error": "INVALID_QUERY", "message": str(exc)}), 400
    except Exception as exc:
        app.logger.exception("Failed to read monitoring metrics")
        return jsonify({"error": "DATABASE_READ_FAILED", "message": str(exc)}), 500


@app.post("/api/events/ai-analysis")
@login_required
def event_ai_analysis():
    body = request.get_json(silent=True)
    event_id = body.get("event_id") if isinstance(body, dict) else None
    if (isinstance(event_id, bool) or not isinstance(event_id, (str, int))
            or not str(event_id).strip() or len(str(event_id)) > 255):
        return jsonify({"error": "INVALID_EVENT_ID"}), 400
    try:
        return jsonify(analyze_event(str(event_id), get_connection, REMEDIATION_ACTIONS))
    except EventNotFound:
        return jsonify({"error": "EVENT_NOT_FOUND"}), 404
    except AnalysisBusy:
        return jsonify({"error": "ANALYSIS_IN_PROGRESS"}), 409
    except Exception:
        # Do not expose event logs, model output, credentials or provider errors.
        app.logger.warning("Event AI analysis unavailable")
        return jsonify({"error": "AI_ANALYSIS_UNAVAILABLE", "message": "AI 분석을 불러오지 못했습니다."}), 503


@app.post("/api/chat")
@login_required
def chat():
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        return (
            jsonify(
                {
                    "error": "OPENAI_API_KEY_MISSING",
                    "message": "backend/.env에 OPENAI_API_KEY를 설정해 주세요.",
                }
            ),
            503,
        )

    data = request.get_json(silent=True) or {}
    message = str(data.get("message") or "").strip()
    if not message:
        return jsonify({"error": "EMPTY_MESSAGE", "message": "질문을 입력해 주세요."}), 400

    raw_history = data.get("history") if isinstance(data.get("history"), list) else []
    event = data.get("event") if isinstance(data.get("event"), dict) else None

    # 토큰/비용 폭증 방지를 위해 최근 대화만 전달한다.
    input_messages = []
    for item in raw_history[-10:]:
        if not isinstance(item, dict):
            continue
        role = "assistant" if item.get("role") == "bot" else "user"
        text = str(item.get("text") or "").strip()
        if text:
            input_messages.append({"role": role, "content": text[:4000]})
    input_messages.append({"role": "user", "content": message[:4000]})

    event_context = "선택된 보안 이벤트 없음"
    if event:
        safe_event = {
            "id": event.get("id"),
            "severity": event.get("severity"),
            "title": event.get("title"),
            "service": event.get("service"),
            "asset": event.get("asset"),
            "detectedAt": event.get("detectedAt"),
            "status": event.get("status"),
            "recommendation": event.get("recommendation"),
            "attackPath": event.get("attackPath"),
            "highlightAssets": event.get("highlightAssets"),
            "details": {
                "attackerIP": (event.get("details") or {}).get("attackerIP"),
                "requestURL": (event.get("details") or {}).get("requestURL"),
                "rule": (event.get("details") or {}).get("rule"),
                "blocked": (event.get("details") or {}).get("blocked"),
                "logs": str((event.get("details") or {}).get("logs") or "")[:4000],
            },
        }
        event_context = json.dumps(safe_event, ensure_ascii=False)

    instructions = f"""
너는 AWS Security Monitoring Center의 보안 분석 어시스턴트다.
답변은 한국어로 간결하고 실무적으로 작성한다.
선택된 이벤트 컨텍스트와 사용자가 제공한 정보에 근거해서만 사실을 단정한다.
정보가 부족하면 부족하다고 명확히 말하고 확인할 AWS 로그/서비스를 제안한다.
실제 AWS 변경이나 차단을 수행했다고 말하지 않는다. 대시보드 사용자가 판단할 수 있도록 분석과 권장 조치만 제공한다.

현재 선택된 이벤트 컨텍스트:
{event_context}
""".strip()

    try:
        client = OpenAI(api_key=api_key)
        response = client.responses.create(
            model=os.getenv("OPENAI_MODEL", "gpt-5.6-luna"),
            instructions=instructions,
            input=input_messages,
            store=False,
            max_output_tokens=int(os.getenv("OPENAI_MAX_OUTPUT_TOKENS", "700")),
        )
        answer = (response.output_text or "").strip()
        if not answer:
            answer = "분석 결과를 생성하지 못했습니다. 다시 질문해 주세요."

        actions = []
        if event:
            if event.get("attackPath"):
                actions.append("공격 경로 강조")
            if event.get("recommendation"):
                actions.append("권장 조치 보기")

        return jsonify({"text": answer, "actions": actions})
    except Exception as exc:
        app.logger.exception("OpenAI chat request failed")
        return (
            jsonify(
                {
                    "error": "OPENAI_REQUEST_FAILED",
                    "message": f"OpenAI API 요청 실패: {exc}",
                }
            ),
            502,
        )


@app.post("/api/ai-diagnosis/run")
@login_required
def run_ai_diagnosis():
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT status FROM ai_diagnosis_runs ORDER BY started_at DESC LIMIT 1"
            )
            latest = cursor.fetchone()
            if latest and latest["status"] in AI_DIAGNOSIS_ACTIVE_STATUSES:
                return (
                    jsonify(
                        {
                            "error": "ALREADY_RUNNING",
                            "message": "이미 진단이 진행 중입니다.",
                        }
                    ),
                    409,
                )

            cursor.execute(
                "INSERT INTO ai_diagnosis_runs "
                "(status, message, requested_by, started_at) "
                "VALUES ('collecting', 'AWS 현재 설정을 수집하는 중입니다...', %s, UTC_TIMESTAMP())",
                (_current_user()["username"],),
            )
            run_id = cursor.lastrowid

    thread = threading.Thread(
        target=_run_ai_diagnosis_job, args=(run_id,), daemon=True
    )
    thread.start()

    return jsonify({"status": "started", "runId": run_id})


@app.get("/api/ai-diagnosis/status")
@login_required
def ai_diagnosis_status():
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT * FROM ai_diagnosis_runs ORDER BY started_at DESC LIMIT 1"
            )
            row = cursor.fetchone()

    if not row:
        return jsonify({"status": "idle"})

    return jsonify(_ai_diagnosis_row_to_json(row))


def _latest_done_ai_diagnosis_run():
    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT * FROM ai_diagnosis_runs WHERE status = 'done' "
                "ORDER BY started_at DESC LIMIT 1"
            )
            return cursor.fetchone()


@app.get("/api/ai-diagnosis/report.xlsx")
@login_required
def ai_diagnosis_report_xlsx():
    row = _latest_done_ai_diagnosis_run()
    if not row or not row.get("result"):
        return (
            jsonify(
                {
                    "error": "NO_COMPLETED_RUN",
                    "message": "완료된 진단 결과가 없습니다. 먼저 진단을 실행해 주세요.",
                }
            ),
            404,
        )

    report = json.loads(row["result"])
    workbook_bytes = _build_ai_diagnosis_workbook(report)
    filename = (
        "ai-diagnosis-"
        f"{_format_datetime(row.get('started_at'), '%Y%m%d-%H%M')}.xlsx"
    )

    return send_file(
        io.BytesIO(workbook_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@app.get("/api/ai-diagnosis/report.pdf")
@login_required
def ai_diagnosis_report_pdf():
    row = _latest_done_ai_diagnosis_run()
    if not row or not row.get("result"):
        return (
            jsonify(
                {
                    "error": "NO_COMPLETED_RUN",
                    "message": "완료된 진단 결과가 없습니다. 먼저 진단을 실행해 주세요.",
                }
            ),
            404,
        )

    report = json.loads(row["result"])
    pdf_bytes = _build_ai_diagnosis_pdf(report)
    filename = (
        "ai-diagnosis-"
        f"{_format_datetime(row.get('started_at'), '%Y%m%d-%H%M')}.pdf"
    )

    return send_file(
        io.BytesIO(pdf_bytes),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=filename,
    )


# 화면(SPA). 해시 라우팅(#/scenario/...)을 쓰므로 서버는 항상 index.html만 주면 된다.
# 위의 /api/* 라우트들이 먼저 매칭되고, 그 외 모든 경로가 마지막으로 이리로 온다.
@app.get("/", defaults={"path": ""})
@app.get("/<path:path>")
def spa(path):
    if app.static_folder and path and os.path.isfile(os.path.join(app.static_folder, path)):
        return app.send_static_file(path)
    return app.send_static_file("index.html")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=True)
